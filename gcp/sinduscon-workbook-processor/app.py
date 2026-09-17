import base64
import html
import os
import re
import shutil
import subprocess
import tempfile
from copy import copy
from pathlib import Path

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

app = FastAPI(title="Sinduscon Workbook Processor")
app.add_middleware(
    CORSMiddleware,
    allow_origins=[],
    allow_origin_regex=r"https://[a-z0-9-]+\.lovable\.app|http://localhost:(5173|8080)",
    allow_methods=["POST", "OPTIONS"],
    allow_headers=["content-type"],
)

MAX_BYTES = 15 * 1024 * 1024
MONTHS = {"JAN": "Janeiro", "FEV": "Fevereiro", "MAR": "Março", "ABR": "Abril", "MAI": "Maio", "JUN": "Junho", "JUL": "Julho", "AGO": "Agosto", "SET": "Setembro", "OUT": "Outubro", "NOV": "Novembro", "DEZ": "Dezembro"}
COLUMN = {"usage": "Uso(s) Alvará", "purpose": "Finalidade", "released": "Área Liberada", "inspection": "Área Vistoria", "residential": "Quantidade de Unidades Residênciais", "non_residential": "Quantidade Unidades Não Residênciais"}

def clean(value):
    return html.unescape(str(value or "")).strip()

def as_number(value):
    if value in (None, ""):
        return 0
    try:
        return float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return 0

def period_from_name(name):
    upper = name.upper()
    token = next((key for key in MONTHS if re.search(rf"(?:_|-){key}(?:_|-)", upper)), None)
    year = re.search(r"20\d{2}", upper)
    if not token or not year:
        raise ValueError("Não foi possível identificar mês e ano no nome do arquivo.")
    return MONTHS[token], int(year.group())

def headers(ws):
    return {clean(ws.cell(1, col).value): col for col in range(1, ws.max_column + 1)}

def require_columns(ws, kind):
    found = headers(ws)
    required = [COLUMN["usage"], COLUMN["purpose"], COLUMN["released"], COLUMN["residential"], COLUMN["non_residential"]]
    if kind == "cvco":
        required.append(COLUMN["inspection"])
    missing = [item for item in required if item not in found]
    if missing:
        raise ValueError("Não encontrei as colunas esperadas: " + ", ".join(missing))
    return found

def commercial(usage): return bool(re.search(r"com[eé]rcio|servi[cç]o", usage, re.I))
def residential_use(usage): return bool(re.search(r"habita[cç][aã]o|residencial", usage, re.I))

def copy_column_dimensions(ws, source, destination):
    src = ws.column_dimensions[get_column_letter(source)]
    dst = ws.column_dimensions[get_column_letter(destination)]
    dst.width, dst.hidden, dst.bestFit = src.width, src.hidden, src.bestFit
    dst.outlineLevel, dst.collapsed = src.outlineLevel, src.collapsed

def insert_columns(ws, index, amount, template):
    max_col = ws.max_column
    dimensions = {col: copy(ws.column_dimensions[get_column_letter(col)]) for col in range(1, max_col + 1)}
    ws.insert_cols(index, amount)
    for old in range(1, max_col + 1):
        new = old + amount if old >= index else old
        source = dimensions[old]
        target = ws.column_dimensions[get_column_letter(new)]
        target.width, target.hidden, target.bestFit = source.width, source.hidden, source.bestFit
        target.outlineLevel, target.collapsed = source.outlineLevel, source.collapsed
    for col in range(index, index + amount):
        copy_column_dimensions(ws, template if template < index else template + amount, col)
        for row in range(1, ws.max_row + 1):
            ws.cell(row, col)._style = copy(ws.cell(row, template if template < index else template + amount)._style)

def move_column(ws, source, destination):
    if source == destination:
        return
    insert_columns(ws, destination, 1, source if source < destination else source + 1)
    actual_source = source + 1 if source >= destination else source
    for row in range(1, ws.max_row + 1):
        original, target = ws.cell(row, actual_source), ws.cell(row, destination)
        target.value, target._style = original.value, copy(original._style)
        if original.number_format: target.number_format = original.number_format
        if original.hyperlink: target._hyperlink = copy(original.hyperlink)
        if original.comment: target.comment = copy(original.comment)
    copy_column_dimensions(ws, actual_source, destination)
    ws.delete_cols(actual_source, 1)

def set_inserted_header(ws, column, title, template):
    source = ws.cell(1, template)
    target = ws.cell(1, column)
    target.value, target._style = title, copy(source._style)
    target.number_format = source.number_format

def normalize_rows(ws, column_map):
    decisions, reviews, removed = [], [], 0
    for row in range(ws.max_row, 1, -1):
        purpose = clean(ws.cell(row, column_map[COLUMN["purpose"]]).value)
        released = as_number(ws.cell(row, column_map[COLUMN["released"]]).value)
        if re.search(r"demoli[cç][aã]o", purpose, re.I) or released == 0:
            ws.delete_rows(row, 1)
            removed += 1
            continue
        usage = clean(ws.cell(row, column_map[COLUMN["usage"]]).value)
        res_cell, non_cell = ws.cell(row, column_map[COLUMN["residential"]]), ws.cell(row, column_map[COLUMN["non_residential"]])
        original_res, original_non = as_number(res_cell.value), as_number(non_cell.value)
        decision = None
        if commercial(usage) and not residential_use(usage) and original_res:
            res_cell.value = None
            decision = "Uso comercial: unidades residenciais foram excluídas; a quantidade não residencial original foi mantida."
        elif residential_use(usage) and not commercial(usage) and original_non:
            non_cell.value = None
            decision = "Uso residencial: unidades não residenciais foram excluídas; a quantidade residencial original foi mantida."
        elif commercial(usage) and residential_use(usage) and original_res and original_non:
            if original_res >= original_non * 3:
                non_cell.value = None
                decision = "Uso misto com predominância residencial: mantidas as unidades residenciais."
            elif original_res <= original_non * 2:
                res_cell.value = None
                decision = "Uso misto comercial: mantidas as unidades não residenciais."
            else:
                reviews.append({"id": f"row-{row}", "rowNumber": row, "usage": usage, "residential": original_res, "nonResidential": original_non, "reason": "Uso misto comercial/residencial em faixa intermediária; revise a classificação."})
        for cell in (res_cell, non_cell):
            if as_number(cell.value) == 0:
                cell.value = None
        if decision:
            decisions.append({"id": f"decision-{row}", "rowNumber": row, "usage": usage, "originalResidential": original_res, "originalNonResidential": original_non, "finalResidential": as_number(res_cell.value), "finalNonResidential": as_number(non_cell.value), "decision": decision})
    return decisions[::-1], reviews[::-1], removed

def prepare_sheet(ws, kind, month, year):
    original_rows = ws.max_row - 1
    initial = require_columns(ws, kind)
    decisions, reviews, removed = normalize_rows(ws, initial)
    insert_columns(ws, 1, 2, 1)
    set_inserted_header(ws, 1, "Mês", 3)
    set_inserted_header(ws, 2, "Ano", 3)
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 1).value, ws.cell(row, 2).value = month, year
    current = require_columns(ws, kind)
    area = current[COLUMN["released"]]
    if kind == "cvco":
        move_column(ws, current[COLUMN["inspection"]], area + 1)
    current = require_columns(ws, kind)
    area = current[COLUMN["released"]]
    insertion = area + (2 if kind == "cvco" else 1)
    insert_columns(ws, insertion, 3, area)
    for offset, title in enumerate(["Área Unidade", "ÁREA RESID", "ÁREA NÃO RESID"]):
        set_inserted_header(ws, insertion + offset, title, area)
    if kind == "cvco":
        current = require_columns(ws, kind)
        move_column(ws, current["Tipo Vistoria"], insertion + 3)
    current = require_columns(ws, kind)
    for row in range(2, ws.max_row + 1):
        area_value = as_number(ws.cell(row, current[COLUMN["released"]]).value)
        res, non = as_number(ws.cell(row, current[COLUMN["residential"]]).value), as_number(ws.cell(row, current[COLUMN["non_residential"]]).value)
        unit = area_value / (res + non) if area_value and (res + non) else None
        ws.cell(row, insertion).value = unit
        ws.cell(row, insertion + 1).value = unit * res if unit and res else None
        ws.cell(row, insertion + 2).value = unit * non if unit and non else None
    return {"rowsRead": original_rows, "rowsRemoved": removed, "rowsKept": ws.max_row - 1, "decisions": decisions, "reviews": reviews}

def convert_to_xlsx(source, workdir):
    if source.suffix.lower() == ".xlsx":
        return source
    subprocess.run(["libreoffice", "--headless", "--convert-to", "xlsx", "--outdir", str(workdir), str(source)], check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=90)
    converted = workdir / f"{source.stem}.xlsx"
    if not converted.exists():
        raise RuntimeError("Não foi possível converter o arquivo XLS mantendo a formatação.")
    return converted

@app.post("/process")
async def process(file: UploadFile = File(...), kind: str = Form(...)):
    if kind not in {"alvaras", "cvco"}:
        raise HTTPException(400, "Tipo de relatório inválido.")
    filename = file.filename or "planilha.xlsx"
    if Path(filename).suffix.lower() not in {".xls", ".xlsx"}:
        raise HTTPException(400, "Envie um arquivo .xls ou .xlsx.")
    payload = await file.read()
    if not payload or len(payload) > MAX_BYTES:
        raise HTTPException(400, "Arquivo vazio ou maior que 15 MB.")
    try:
        month, year = period_from_name(filename)
        with tempfile.TemporaryDirectory() as raw_dir:
            workdir = Path(raw_dir)
            source = workdir / filename
            source.write_bytes(payload)
            converted = convert_to_xlsx(source, workdir)
            workbook = load_workbook(converted)
            worksheet = workbook[workbook.sheetnames[0]]
            report = prepare_sheet(worksheet, kind, month, year)
            output_name = f"{'Alvaras' if kind == 'alvaras' else 'CVCO'}_tratado_{month}_{year}.xlsx"
            output = workdir / output_name
            workbook.save(output)
            return {"fileName": output_name, "contentBase64": base64.b64encode(output.read_bytes()).decode(), "report": {"kind": kind, "fileName": filename, "month": month, "year": year, **report}}
    except ValueError as error:
        raise HTTPException(400, str(error)) from error
    except Exception:
        raise HTTPException(500, "Não foi possível tratar a planilha preservando sua formatação.")
