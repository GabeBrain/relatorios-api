#!/usr/bin/env python
"""Extract a machine-readable validation reference from the official FIERGS PPTX.

The output intentionally preserves the source structure instead of guessing metric
semantics: slide text, native PowerPoint tables and native chart series. Raster-only
slides are marked so their values can be completed during visual validation.

Usage:
    python scripts/extract-fiergs-reference.py SOURCE.pptx OUTPUT.json
"""

from __future__ import annotations

import argparse
import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

from pptx import Presentation


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def clean_text(value: str) -> str:
    return " ".join(value.replace("\u00a0", " ").split())


def chart_data(shape: Any) -> dict[str, Any]:
    chart = shape.chart
    plots = []
    for plot_index, plot in enumerate(chart.plots, start=1):
        categories = []
        try:
            categories = [json_value(category.label) for category in plot.categories]
        except (AttributeError, TypeError, ValueError):
            pass

        series = []
        for item in plot.series:
            values = []
            try:
                values = [json_value(value) for value in item.values]
            except (AttributeError, TypeError, ValueError):
                pass
            series.append({"name": clean_text(str(item.name or "")), "values": values})
        plots.append({"index": plot_index, "categories": categories, "series": series})

    title = None
    try:
        if chart.has_title:
            title = clean_text(chart.chart_title.text_frame.text)
    except (AttributeError, ValueError):
        pass
    return {"title": title, "plots": plots}


def table_data(shape: Any) -> list[list[str]]:
    return [[clean_text(cell.text) for cell in row.cells] for row in shape.table.rows]


def extract(source: Path) -> dict[str, Any]:
    presentation = Presentation(source)
    slides = []
    for number, slide in enumerate(presentation.slides, start=1):
        texts = []
        tables = []
        charts = []
        pictures = 0
        for shape in slide.shapes:
            text = clean_text(getattr(shape, "text", ""))
            if text:
                texts.append(text)
            if getattr(shape, "has_table", False):
                tables.append(table_data(shape))
            if getattr(shape, "has_chart", False):
                charts.append(chart_data(shape))
            if getattr(shape, "shape_type", None) == 13:  # MSO_SHAPE_TYPE.PICTURE
                pictures += 1
        slides.append(
            {
                "slide": number,
                "texts": texts,
                "tables": tables,
                "charts": charts,
                "pictures": pictures,
                "validationMode": "numeric" if tables or charts else "visual",
            }
        )
    return {
        "schemaVersion": 1,
        "source": source.name,
        "slideCount": len(slides),
        "slides": slides,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    result = extract(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    numeric = sum(slide["validationMode"] == "numeric" for slide in result["slides"])
    print(f"{result['slideCount']} slides; {numeric} with native numeric evidence; output: {args.output}")


if __name__ == "__main__":
    main()
