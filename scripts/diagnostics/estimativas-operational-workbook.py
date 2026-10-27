"""Internal candidate workbook, never an approved Belogic import file.
The model results are immutable outputs from the versioned JS reference, not Excel recalibration.
"""
import json
from pathlib import Path
import xlsxwriter
from openpyxl import load_workbook

ROOT = Path('.local-inputs/estimativas-vendas/operacional-20261027')
data = json.loads((ROOT / 'amostra.json').read_text(encoding='utf-8'))
target = ROOT / 'CANDIDATO_NAO_IMPORTAR_2026-10-27.xlsx'
wb = xlsxwriter.Workbook(target, {'strings_to_formulas': False, 'strings_to_urls': False})
base = {'font_name': 'Arial', 'font_size': 10, 'valign': 'vcenter'}
text = wb.add_format(base)
integer = wb.add_format({**base, 'num_format': '#,##0'})
header = wb.add_format({**base, 'bold': True, 'bg_color': '#4B5930', 'font_color': '#FFFFFF', 'text_wrap': True})
title = wb.add_format({**base, 'font_size': 14, 'bold': True})
warning = wb.add_format({**base, 'font_color': '#9C4100', 'bold': True})
wrap = wb.add_format({**base, 'text_wrap': True})

def sheet(name, headings, widths):
    s = wb.add_worksheet(name)
    s.hide_gridlines(2)
    s.write(0, 0, 'CANDIDATO INTERNO — NÃO IMPORTAR', title)
    s.write(1, 0, '27/10/2026 · corte set/2026 · fonte congelada, sem confirmação da coleta atual', text)
    s.write(2, 0, 'Layout, granularidade, preço, cadência e alternância aguardam aceite.', warning)
    for col, (label, width) in enumerate(zip(headings, widths)):
        s.write(4, col, label, header)
        s.set_column(col, col, width, text)
    s.set_row(4, 32)
    s.freeze_panes(5, 2)
    return s

s = sheet('Estimativa candidata', ['Empreendimento ID', 'Tipologia ID', 'Nome', 'Cidade', 'Tipo / padrão', 'Origem', 'Início alvo', 'Fim alvo', 'Cadência (meses)', 'Unidades', 'Estoque inicial', 'Vendido acumulado inicial', 'Vendido acumulado projetado', 'Venda líquida do período', 'Estoque final', 'Método / versão', 'Estado'], [18,18,34,20,34,13,13,13,16,12,16,19,21,19,16,42,28])
for i, r in enumerate(data['sample'], 5):
    n = i + 1
    vals = [r['buildingId'], r['typologyId'], r['name'], r['city'], r['type']+' / '+r['pattern'], r['cutoff'], r['periodStart'], r['competence'], r['cadenceMonths'], r['units'], r['initialStock'], r['initialCumulative'], r['estimatedCumulative'], r['estimatedNetPeriod'], r['finalStock'], r['method']['id']+' / '+r['method']['version'], r['status']]
    for col, v in enumerate(vals): s.write(i, col, v, integer if isinstance(v, int) else text)
    s.write_formula(i,11, f'=J{n}-K{n}',integer,r['initialCumulative'])
    s.write_formula(i,13, f'=M{n}-L{n}',integer,r['estimatedNetPeriod'])
    s.write_formula(i,14, f'=J{n}-M{n}',integer,r['finalStock'])
s.autofilter(4,0,4+len(data['sample']),16)
notes_row=7+len(data['sample'])
s.write(notes_row,0,'Sem preço, venda bruta ou distratos projetados; ausência não significa zero.',warning)
s.write(notes_row+1,0,'Acumulado projetado é saída congelada do motor. Alterar células não recalcula Finotti.',text)
s.write(notes_row+2,0,'Intenção: estimation=true na competência-alvo; base_id e ivv_batch_id indefinidos.',text)
s.write(notes_row+3,0,'Origem, modelo e linhas-fonte estão no manifesto e no histórico acompanhante.',text)

chosen=[]
for r in data['selection']:
    if r['calculationAllowed']: chosen.append(r)
for reason in sorted(data['reasonCounts']):
    r=next(r for r in data['selection'] if reason in r['reasons'])
    if r not in chosen: chosen.append(r)
s=sheet('Seleção revisável', ['Empreendimento ID','Cidade','Tipo','Cadência hipótese','Estoque total','Decisão','Motivos','Tipologias no corte','Fonte'],[19,23,15,19,17,26,95,34,27])
for i,r in enumerate(chosen,5):
    vals=[r['buildingId'],r['city'],r['type'],r['cadence'],r['stockBuilding'],r['decision'],'; '.join(r['reasons']) or 'Passou no filtro candidato; pendências externas mantidas',', '.join(r['typologies']),r['source']]
    for c,v in enumerate(vals):s.write(i,c,v,wrap if c==6 else integer if isinstance(v,int) else text)
    s.set_row(i,55)
s.autofilter(4,0,4+len(chosen),8)

s=sheet('Histórico da amostra',['Empreendimento ID','Tipologia ID','Competência','Unidades','Estoque','Venda líquida reportada','Foi estimado?','Lançamento','Status','Linha de origem'],[19,18,15,13,13,24,17,15,16,62])
for i,r in enumerate(data['inputs'],5):
    vals=[r['buildingId'],r['typologyId'],r['period'],r['units'],r['stock'],r['net'],'Sim' if r['estimated'] else 'Não',r['release'],r['status'],r['sourceRow']]
    for c,v in enumerate(vals):s.write(i,c,v,integer if isinstance(v,int) else text)
s.autofilter(4,0,4+len(data['inputs']),9)

s=sheet('Manifesto',['Item','Valor'],[37,120])
manifest=[('Finalidade','Revisão interna; não homologado; nenhum registro importável'),('Recorte','Três arquivos IVV: horizontal, Cuiabá e Belo Horizonte; posição setembro/2026'),('Cadência','Hipóteses mensal e trimestral avaliadas separadamente; não dividir venda trimestral por três'),('Quantidade de empreendimentos',data['uniqueBuildings']),('Avaliações (2 cadências por ID)',len(data['selection'])),('Candidatos numéricos',data['sampleSize']),('Lote importável',0),('Política','4 posições consecutivas não estimadas; estoque total 6–30; oferta estável; uma tipologia; sem registro-alvo'),('Limites','Nenhuma divisão da metade ou rateio de tipologias; sem confirmação da ausência de coleta em outubro'),('Reconciliação','Conservar estimativa e coleta separadas; destino da diferença e revisão histórica aguardam decisão'),('Arredondamento','Motor original arredonda acumulado; fluxo = acumulado final - inicial; sem segundo arredondamento'),('Execução','node scripts/diagnostics/estimativas-operational-sample.mjs; python scripts/diagnostics/estimativas-operational-workbook.py'),('Seleção integral','amostra.json ao lado deste arquivo; critérios e motivos por ID/cadência'),('Fonte externa','https://drive.google.com/drive/folders/1p9bDnqLdfgBfvgDovIT4Nvp5lnzNCeP6')]
manifest += [(x['path'],x['sha256']) for x in data['sources']]
manifest += [('Modelo '+x['file'],x['sha256']) for x in data['models']]
for i,(k,v) in enumerate(manifest,5):s.write(i,0,k,wrap);s.write(i,1,v,wrap);s.set_row(i,32)
wb.close()

# Readback checks with both stored numeric values and visible formulas.
cached=load_workbook(target,data_only=True)
formulas=load_workbook(target,data_only=False)
for i,r in enumerate(data['sample'],6):
    s=cached['Estimativa candidata']
    assert s.cell(i,14).value == r['estimatedNetPeriod']
    assert s.cell(i,15).value == r['finalStock']
    assert s.cell(i,11).value-s.cell(i,14).value==s.cell(i,15).value
    assert formulas['Estimativa candidata'].cell(i,14).data_type=='f'
assert len(cached.sheetnames)==4
assert len({(r['buildingId'],r['typologyId'],r['competence']) for r in data['sample']})==len(data['sample'])
print(json.dumps({'file':str(target),'sample':len(data['sample']),'selectionExamples':len(chosen),'tabs':cached.sheetnames,'readback':'passed'},ensure_ascii=True))
